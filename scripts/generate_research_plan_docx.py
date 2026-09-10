#!/usr/bin/env python3
"""Generate the Adaptive FAST-LIO2 research and experiment plan as a DOCX."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "paper" / "Adaptive_FAST_LIO2_科研计划与实验方案.docx"


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False, color=None, size=9.0, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(document, headers, rows, widths=None, font_size=8.6):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for idx, value in enumerate(headers):
        set_cell_text(
            header.cells[idx], value, bold=True, color=(255, 255, 255),
            size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(header.cells[idx], "1F4E78")
        if widths:
            header.cells[idx].width = Cm(widths[idx])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=font_size)
            if row_index % 2:
                set_cell_shading(cells[idx], "EAF2F8")
            if widths:
                cells[idx].width = Cm(widths[idx])
            set_cell_margins(cells[idx])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(document, text, *, bold=False, color=None, align=None):
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold, color=color)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_code(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.right_indent = Cm(0.8)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    run = paragraph.add_run(lines)
    set_run_font(run, name="Consolas", size=8.5)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end_run = paragraph.add_run(" 页")
    set_run_font(end_run, size=9)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    heading_settings = {
        "Title": (22, "1F4E78"),
        "Heading 1": (16, "1F4E78"),
        "Heading 2": (13, "2F5597"),
        "Heading 3": (11.5, "404040"),
    }
    for style_name, (size, color) in heading_settings.items():
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Adaptive FAST-LIO2｜科研计划与实验执行方案")
    set_run_font(run, size=8.5, color=(100, 100, 100))
    add_page_number(section.footer.paragraphs[0])


def add_status_box(document):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "D9EAF7")
    set_cell_text(
        cell,
        "当前阶段：算法原型完成后的稳定化、消融验证与版本冻结阶段\n"
        "当前优先工作：审核已有证据、核对基线公平性、补齐核心消融",
        bold=True,
        color=(31, 78, 121),
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    document.add_paragraph()


def build_document():
    document = Document()
    configure_document(document)
    document.core_properties.title = "Adaptive FAST-LIO2 科研计划与实验执行方案"
    document.core_properties.subject = "退化感知地图管理与后端优化研究计划"
    document.core_properties.author = "Adaptive FAST-LIO2 Project"
    document.core_properties.keywords = "FAST-LIO2, degeneracy, adaptive map, SLAM, research plan"

    # Cover page
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("Adaptive FAST-LIO2")
    set_run_font(run, size=26, bold=True, color=(31, 78, 121))
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("退化感知 SLAM 科研计划与实验执行方案")
    set_run_font(run, size=20, bold=True, color=(64, 64, 64))

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cover_rows = [
        ("项目目录", "/home/romi/Adaptive_FAST_LIO2"),
        ("当前候选逻辑", "基于提交 992f393 的 Adaptive Map v1"),
        ("当前研究阶段", "稳定化、消融验证与版本冻结"),
        ("文档日期", date.today().isoformat() + "｜修订版 v2"),
    ]
    for row, values in zip(table.rows, cover_rows):
        set_cell_text(row.cells[0], values[0], bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(row.cells[0], "D9EAF7")
        set_cell_text(row.cells[1], values[1], size=10)
    document.add_page_break()

    document.add_heading("1. 当前阶段与研究定位", level=1)
    add_status_box(document)
    add_body(
        document,
        "当前项目已经完成基础系统搭建、主要数据集转换、真值评估流程和 Adaptive Map 原型开发，"
        "并在 GEODE、SubT 与 NTU 等数据上进行了多轮初步实验。当前工作的重点不再是继续叠加功能，"
        "而是先审核已有结果，明确可复用、需重评分和需补跑的记录，再通过受控实验验证模块贡献。"
        "当前版本是候选版本，稳定性与有效性仍须验证，不能由回退动作直接推定。",
    )
    add_body(
        document,
        "当前前端已回退到基于 992f393 的单文件候选逻辑；Transition Guard、历史可靠帧快照和弱方向软权重"
        "均不属于当前代码。代码仍保留窗口累计 yaw 变化，用于区分拐角与长直持续退化，但它只是分类条件，"
        "不是主动掉头修正。",
    )

    document.add_heading("2. 核心研究问题与假设", level=1)
    add_body(document, "以下均为待检验假设，不是已证明结论。论文主线是根据匹配质量和退化状态管理地图更新；新颖性需要通过与相关工作的具体机制比较来论证。")
    hypotheses = [
        ("H1", "低质量匹配点持续写入局部地图会造成地图污染，并进一步放大后续配准误差。"),
        ("H2", "与单帧驱动的入图策略相比，时间窗口策略能减少不合理切换，并改善持续退化段的定位表现。"),
        ("H3", "在入图数量接近时，按几何方向选择点优于不按方向选择点；收益不仅来自减少点数。"),
        ("H4（可选）", "在固定前端输入下，可靠的后端约束能改善全局一致性；方向性噪声与连续回环验证作为候选扩展，按诊断证据决定是否实现。"),
    ]
    for key, text in hypotheses:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(f"{key}：")
        set_run_font(run, bold=True, color=(31, 78, 121))
        run = paragraph.add_run(text)
        set_run_font(run)

    document.add_heading("3. 完整科研流程计划", level=1)
    plan_rows = [
        ("1", "证据整理", "审核已有版本、配置、bag、评分和完整性", "立即执行", "明确可用、重评分、补跑清单", "实验审计表"),
        ("2", "公平性检查", "核对预处理、更新顺序、参数和时间戳", "待核对", "基线差异能够解释", "实现与协议差异表"),
        ("3", "核心机制验证", "质量筛选、单帧/窗口、配额及点数控制", "已有探索结果", "明确策略贡献和适用条件", "受控消融结果"),
        ("4", "最小必要修正", "只修改已定位的失败原因", "按证据启动", "收益可重复，无明显跨场景退步", "候选版本与回归报告"),
        ("5", "冻结与泛化", "固定参数，测试其他序列", "待核心验证", "报告优势、代价和失败边界", "冻结版本与泛化结果"),
        ("6", "后端评估", "固定前端，先比较现有后端开关", "原型已有", "判断独立贡献是否有证据支持", "后端评价与扩展决策"),
        ("7", "统计与写作", "补齐主表、开销、失败案例与论文", "写作可并行", "每个结论有可追溯证据", "论文及复现材料"),
    ]
    add_table(
        document,
        ["阶段", "名称", "核心任务", "状态", "完成标准", "产物"],
        plan_rows,
        widths=[0.8, 2.0, 4.1, 2.0, 4.1, 3.0],
        font_size=8.0,
    )

    document.add_heading("4. 当前阶段的具体工作", level=1)
    document.add_heading("4.1 先审核已有实验，再固定候选版本", level=2)
    add_table(document, ["审核分类", "判定条件", "处理方式"], [
        ("可直接使用", "版本、参数、bag 和评估协议可追溯且符合目标比较", "登记原路径；无需重跑或复制"),
        ("需要重新评分", "原始轨迹完整，主要差异在坐标/时间关联或评分口径", "保留原始轨迹，按冻结协议重评分"),
        ("需要补跑", "目标版本或实验组缺失，配置不一致，记录不足", "说明补跑理由，分配唯一运行编号"),
        ("仅作历史参考", "源码或条件无法确认，无法公平比较", "保留并标注不纳入当前主表"),
    ], widths=[3.0, 7.0, 6.0])
    add_body(
        document,
        "将当前基于 992f393 的实现整理为 Adaptive Map v1 候选版本。提交时只纳入前端、后端、配置和必要评估脚本，"
        "不得使用 git add . 将历史实验文件的删除或未跟踪数据混入源码提交。",
    )
    add_bullet(document, "固定代码提交、配置文件和编译方式。")
    add_bullet(document, "固定初始化帧数、外参、IMU 噪声、体素大小、探测范围和播放速度。")
    add_bullet(document, "固定每个数据集的时间戳定义、允许时间差和坐标变换流程。")
    add_bullet(document, "验证通过后再创建 adaptive-map-v1-stable 标签。")

    document.add_heading("4.2 基线公平性检查", level=2)
    equivalence_rows = [
        ("FAST-LIO2", "未经 Adaptive 模块修改的独立基线", "建立原始性能参照"),
        ("Ours w/o Adaptive", "使用本项目框架但关闭 Adaptive Map 和窗口", "排除工程实现差异"),
        ("Adaptive Map", "开启质量筛选、状态机和方向/体素策略", "验证创新点1整体作用"),
    ]
    add_table(document, ["方法", "配置含义", "验证目的"], equivalence_rows, widths=[3.0, 7.0, 6.0])
    add_body(
        document,
        "如果 Ours w/o Adaptive 与 FAST-LIO2 的差异无法由时间戳、预处理或日志采样解释，应暂停算法结论，"
        "优先排查实现公平性。关闭模块不等于工程等价，ATE 接近也不能证明实现相同。"
        "先对照预处理、去畸变、滤波更新顺序、地图管理和日志采样，再比较相同时间支持下的轨迹。",
    )

    document.add_heading("4.3 代表性开发序列：按缺口补实验", level=2)
    regression_rows = [
        ("GEODE Tunnel5", "典型隧道、具备官方真值", "检查基础稳定性和 GEODE 官方评分流程"),
        ("SubT Hawkins Long Corridor", "长走廊持续弱约束", "验证 Persistent 状态和总入图配额"),
        ("NTU SPMS2", "掉头区域容易发生瞬态配准失效", "判断是否确实需要 Transition Guard"),
    ]
    add_table(document, ["序列", "场景特征", "验证目标"], regression_rows, widths=[4.1, 5.8, 6.1])
    add_body(document, "上述序列已多次参与开发，不要求全部重新运行。审核后选择一个序列补最关键的受控比较。GEODE 未触发 Persistent 时，只能支持质量筛选和 Transient 策略的结论；不得为增加触发次数而针对测试序列降低阈值。")

    document.add_heading("5. Innovation 1：Adaptive Map 消融计划", level=1)
    ablation_rows = [
        ("A", "FAST-LIO2", "固定基线实现", "与外部基线比较"),
        ("B", "Ours w/o Adaptive", "关闭自适应地图模块和窗口策略", "检查工程框架差异"),
        ("C", "基础质量筛选", "仅质量门限；不启用退化配额", "质量筛选的贡献"),
        ("D", "单帧自适应策略", "单帧退化判断驱动方向/新区域配额", "单帧入图策略的贡献"),
        ("E", "完整 Adaptive Map", "在 D 基础上加入窗口与 Persistent 策略", "时序策略的增量贡献"),
        ("F", "点数控制组", "入图预算接近目标策略，但不按法向方向选择", "区分几何选择与单纯减点"),
    ]
    add_table(
        document,
        ["编号", "方法", "实际行为", "回答的问题"],
        ablation_rows,
        widths=[0.9, 3.6, 7.0, 4.5],
        font_size=8.2,
    )

    add_body(document, "所有核心消融关闭后端。表中是实验设计目标，执行前须检查代码是否真正实现这些行为。仅增加状态分类但不改变入图行为，不构成有独立轨迹作用的消融。Persistent 总配额可先在 SubT 上单独比较开启和关闭。")
    add_body(document, "F 与目标组应共享基础质量条件、体素处理和预算规则，只改变候选点的选择方式；选择规则预先固定，随机选择时登记种子。记录实际入图数量、地图大小和计算开销，量化剩余数量差异。不得根据真值误差选择点或预算。")

    document.add_heading("5.1 Transition Guard 的决策原则", level=2)
    add_body(
        document,
        "Transition Guard 当前不在代码中，不应直接作为正式方法组成部分。只有在同版本 SPMS2 配对实验中，"
        "反复证明掉头阶段会污染地图并造成永久发散时，才重新实现。",
    )
    guard_rules = [
        "先核对事件顺序：转弯、匹配质量变化、位姿误差、可疑点入图、后续误差增长；时序相关不能单独证明因果。",
        "检查去畸变、配准与同步问题；若初始错误先于入图，Guard 只能限制传播，不能宣称修复最初的错误位姿。",
        "在当前帧正常 ESIKF 更新之后判断，不改变 FAST-LIO2 基础状态更新。",
        "只暂停或限制可疑点入图，不对整帧测量统一降权。",
        "设置最大持续帧数，避免长期冻结地图。",
        "采用进入/退出滞回和重新武装条件，避免连续重复触发。",
        "必须进行相同提交、相同配置下的 Guard-off / Guard-on 配对实验。",
    ]
    for rule in guard_rules:
        add_bullet(document, rule)

    document.add_heading("6. 后端评估与可选扩展", level=1)
    add_body(document, "先固定前端，比较后端关闭与现有后端开启的结果，并检查接受回环的可靠性。方向软权重、各向异性噪声和连续回环验证不是必做任务；只有诊断支持且预期收益明确时，才逐项实现。若后端证据不足，可保留系统组件身份，不强行作为独立创新贡献。")
    backend_rows = [
        ("6.1", "固定前端输入", "防止前端版本变化干扰后端比较"),
        ("6.2", "保留标量噪声控制组", "建立当前后端基线"),
        ("可选", "各向异性里程计协方差", "方向信息可靠且标量权重存在问题时验证"),
        ("可选", "连续候选一致性", "存在偶然错误回环证据时验证"),
        ("可选", "回环前后局部一致性检查", "检验错误回环对全局轨迹的影响"),
        ("6.6", "无可靠回环时维持正常里程计边", "避免没有依据地削弱全部约束"),
    ]
    add_table(document, ["步骤", "内容", "目的"], backend_rows, widths=[1.8, 7.0, 7.0])
    add_body(
        document,
        "后端属于松耦合优化：它根据前端里程计建立关键帧与位姿图，不回写前端状态和 iKD-Tree。"
        "因此后端能够修正全局轨迹，但不能从根本上解决前端掉头瞬间的错误匹配。",
    )

    document.add_heading("7. 数据范围与开发/泛化划分", level=1)
    dataset_rows = [
        ("GEODE Offroad", "1–3", "非结构化、颠簸环境的泛化性能", "官方 Gamma GNSS 转换与 rmse.py"),
        ("GEODE Tunneling", "1–5", "隧道退化主实验", "官方 Gamma Leica 转换与 rmse.py"),
        ("GEODE Waterway", "Short、Medium", "开放水面和弱几何结构", "官方 Gamma GNSS 转换与 rmse.py"),
        ("SubT", "Hawkins Long Corridor", "Persistent 持续退化证据", "evo_ape，固定时间差"),
        ("NTU VIRAL", "SPMS1–3", "跨雷达与掉头场景泛化", "棱镜补偿、插值和 SE(3) 对齐"),
        ("Corridor02", "单序列", "地图和轨迹定性展示", "无可靠真值，不进入定量主表"),
    ]
    add_table(document, ["数据集", "序列", "用途", "评估方式"], dataset_rows, widths=[3.2, 3.1, 5.3, 4.7], font_size=8.3)
    add_body(document, "此表为候选数据范围，不代表必须全量重跑。SPMS2、Tunnel2、Tunnel5、SubT 已反复参与调参，应登记为开发序列。其他序列先审核使用历史；只有未参与方法与参数选择的序列才能作为独立泛化验证。若所有序列均参与过开发，应如实说明并寻找额外未使用序列。")
    add_body(document, "评分流程须对照本地官方脚本与说明核验，记录脚本版本、棱镜/坐标补偿、插值、对齐和时间关联参数。额外采用的时间偏移或容差不得自动称为官方规定。")

    document.add_heading("8. 统一实验与评分规范", level=1)
    protocol_rows = [
        ("主指标", "ATE RMSE；不同真值来源的数据集不混合计算总体平均值"),
        ("辅助指标", "RPE、最大误差、最终误差、成功率、运行时间和内存"),
        ("重复实验", "以 3 次配对运行为初步重复性检查；报告各次值、均值、样本标准差和成功率。按波动与效果大小决定是否扩充，并对各方法采用同一停止规则"),
        ("播放速度", "正式实验统一 1.0 倍；不得为某个方法单独降低速度"),
        ("时间戳", "按位姿物理时刻确定帧首/帧尾及偏移，记录正负号和补偿顺序；不得按最小 ATE 选择 offset，也不得重复补偿"),
        ("参数", "参数冻结后禁止针对单条测试序列调整"),
        ("RViz", "精度实验统一开启或统一关闭；计算性能实验应关闭"),
        ("失败处理", "分开记录算法失败、输入/配置错误、运行中断和未完成；适配失败不能直接等同算法失效。保留失败数据，成功运行 ATE 与成功率分别报告"),
        ("覆盖率", "记录完整运行时长、关联点数、有效真值覆盖率和轨迹缺口；避免短轨迹 ATE 被误当完整序列性能"),
        ("显著性", "三次运行不能保证统计显著；效果与波动接近时应报告不确定性，不能强称有效"),
    ]
    add_table(document, ["项目", "统一要求"], protocol_rows, widths=[3.2, 13.2])

    document.add_heading("9. 实验数据组织规范", level=1)
    add_code(
        document,
        "experiments/adaptive_map_v1/\n"
        "├── README.md\n"
        "├── experiment_registry.csv\n"
        "├── audit.csv\n"
        "├── validation/\n"
        "│   ├── geode_tunnel5/{fastlio2,ours_no_adaptive,adaptive_map}/run01/\n"
        "│   ├── subt_hawkins/\n"
        "│   └── ntu_spms2/\n"
        "├── formal/{geode,subt,ntu}/\n"
        "└── summary/{validation_ate.csv,formal_ate.csv,figures/}",
    )
    add_body(document, "每次运行固定保存以下文件：")
    files_rows = [
        ("runtime.csv", "逐帧前端指标、状态和入图统计"),
        ("trajectory.tum", "用于定量评估的轨迹"),
        ("metrics.txt", "ATE、RPE、最大误差等结果"),
        ("launch_command.txt", "实际启动与播放命令"),
        ("config.yaml", "配置快照"),
        ("git_commit.txt", "源码提交哈希；若有未提交修改，附 diff 快照与构建记录"),
        ("manifest.yaml", "bag 校验值、有效参数与命令行覆盖、评估脚本版本、时间定义和覆盖率"),
        ("notes.md", "RViz 现象、异常日志和失败说明"),
    ]
    add_table(document, ["文件", "作用"], files_rows, widths=[4.3, 12.0])
    add_body(
        document,
        "旧实验保留原路径，通过登记表引用，避免搬迁或重复复制。validation 用于开发；formal 汇总符合冻结协议的结果，"
        "审核合格的旧结果可通过索引引用，不因目录年代而强制重跑。一次运行一个唯一编号；输出文件名差异由清单映射。",
    )
    add_body(document, "审计表每行记录：运行 ID、数据集/序列、方法、源码版本、配置及覆盖参数、bag 版本、评分版本、时间协议、运行完整性、开发/测试属性、原始路径、审核分类和补跑理由。未知项明确标记未知，不从文件夹名称推定。")

    document.add_heading("10. 论文图表与证据链", level=1)
    figure_items = [
        "系统整体框架图：前端 Adaptive Map 与松耦合后端的输入输出关系。",
        "Normal / Transient / Persistent 三状态切换图。",
        "基础地图质量筛选、方向分箱和体素选择流程图。",
        "代表序列中的退化指标、状态、入图率和拒绝原因时间序列。",
        "FAST-LIO2、Ours w/o Adaptive、Adaptive Map 与 Ours-Full 的 ATE 主表。",
        "质量筛选、单帧/窗口、配额及近似等点数控制的消融表；Guard 有效后才纳入。",
        "Tunnel、SubT 和 SPMS 的代表性轨迹与局部地图对比。",
        "运行时间、内存、成功率以及失败案例分析。",
    ]
    for item in figure_items:
        add_number(document, item)

    document.add_heading("11. 时间预算与阶段验收", level=1)
    add_body(document, "以下为可调整预算，不是十周内完成全部扩展的承诺。以验收条件决定是否推进；条件不满足时优先定位原因或缩小方法范围。写作从现在开始并行整理问题、公式和协议。")
    schedule_rows = [
        ("2–3 天", "审核现有实验和评分协议", "确定可用、重评分和补跑清单"),
        ("3–5 天", "检查基线实现、参数与时间语义", "差异有证据解释"),
        ("1–2 周", "补齐代表序列核心消融及点数控制", "明确贡献、波动与适用范围"),
        ("按需 1–2 周", "最小修正；有证据才恢复 Guard", "收益可重复且无明显回归"),
        ("约 1 周", "冻结前端并开展泛化验证", "参数固定，开发/测试属性清楚"),
        ("按需 1–2 周", "现有后端评价与必要扩展", "明确是否支持独立贡献"),
        ("1–2 周", "补齐正式结果、开销和失败分析", "结果可追溯，报告完整"),
        ("并行；收尾约 1 周", "方法写作、图表整理与复现检查", "论点与证据一致"),
    ]
    add_table(document, ["时间", "主要工作", "里程碑"], schedule_rows, widths=[2.0, 9.2, 5.2])

    document.add_heading("12. 近期执行顺序", level=1)
    next_steps = [
        "先建立审计表，引用现有结果原路径，不立即启动完整矩阵。",
        "逐项核对代码、有效参数、bag、时间协议、评估脚本和运行完整性，分类为可用、重评分、补跑或历史参考。",
        "记录当前候选源码与基线差异；构建编译由用户执行并记录版本，不因尚未验证就命名为稳定版。",
        "选择一个代表序列补最关键的消融，每次运行先写明唯一研究问题、控制变量和比较对象。",
        "根据实际入图行为比较质量筛选、单帧/窗口策略及近似等点数控制；Persistent 仅在实际触发序列上讨论。",
        "证据足够后冻结前端；Guard 与后端扩展各自按诊断结果决定是否开展。",
    ]
    for step in next_steps:
        add_number(document, step)

    document.add_heading("13. 阶段退出标准", level=1)
    exit_rows = [
        ("代码可追溯", "源码、配置、评估脚本均对应唯一 Git commit"),
        ("基线公平", "Ours w/o Adaptive 与 FAST-LIO2 的差异已解释"),
        ("运行稳定", "代表序列不出现由新增逻辑引入的系统性发散"),
        ("机制有效", "日志能够证明状态切换、筛选和配额实际执行"),
        ("效果可重复", "报告效应大小与重复波动；证据不足时标注不确定或增加配对运行，不选择最优值"),
        ("机制可解释", "点数控制能区分几何选择与减点效应；未触发策略不宣称贡献"),
        ("结论边界", "明确开发序列、泛化验证、失败类别及不适用场景；不要求所有场景都获益"),
        ("参数冻结", "正式实验前不再针对单条测试序列调参"),
    ]
    add_table(document, ["检查项", "退出要求"], exit_rows, widths=[3.8, 12.5])

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(10)
    conclusion.paragraph_format.space_after = Pt(0)
    conclusion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = conclusion.add_run(
        "当前首要任务：审核已有证据，再用最少的受控实验补齐 Adaptive Map 的贡献证据。"
    )
    set_run_font(run, size=12, bold=True, color=(192, 80, 77))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
